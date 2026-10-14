"""Historical costs, current state and time-bounded dependency attribution."""
from decimal import Decimal
from io import BytesIO
from unittest import TestCase
from unittest.mock import patch

from openpyxl import load_workbook

from aws_inventory.estimates import month_windows
from aws_inventory.inventory_compute import collect_ec2
from aws_inventory.inventory_extended import collect_emr
from aws_inventory.service_exports import build_service_workbook
from tests.test_collectors import _context
from tests.test_html_report import actual, build, quote, resource


class ResourceLifetimeTests(TestCase):
    def test_emr_creation_from_sdk_bounds_dependency_costs(self):
        for source in ('description', 'summary', 'missing'):
            with self.subTest(source=source):
                ctx = _context({'emr'})
                summary = {'Id': 'j-cluster', 'Name': 'Analytics', 'Status': {'State': 'RUNNING'}}
                cluster = {**summary, 'Status': dict(summary['Status'])}
                if source != 'missing':
                    target = cluster if source == 'description' else summary
                    target['Status']['Timeline'] = {'CreationDateTime': '2026-09-16',
                                                    'ReadyDateTime': '2026-09-17'}
                def pages(service, operation, region, result_key, **kwargs):
                    if operation == 'list_clusters':
                        return [summary]
                    if operation == 'list_instances':
                        return [{'Ec2InstanceId': 'i-one'}]
                    self.fail(f'Unexpected operation: {operation}')
                with patch.object(ctx, 'safe_pages', side_effect=pages), \
                        patch.object(ctx, 'safe_call', return_value={'Cluster': cluster}):
                    collect_emr(ctx, 'eu-west-1')
                model = build([*ctx.rows, resource(created_at='2026-01-01')],
                              period_start='2026-08-01')
                parent = next(r for r in model['resources'] if r['service'] == 'emr')
                worker = next(r for r in model['resources'] if r['service'] == 'ec2')
                if source == 'missing':
                    self.assertFalse(parent['lifetime']['creation_known'])
                else:
                    self.assertEqual(parent['lifetime']['start'], '2026-09-16T00:00:00+00:00')
                    self.assertEqual(parent['totals']['2026-08']['amounts']['USD'], 0)
                    self.assertEqual(parent['totals']['2026-09']['components'][worker['key']]['amount'], Decimal('43.2'))
                    self.assertEqual(worker['months']['2026-09']['amount'], Decimal('86.4'))

    def test_historical_stopped_state_does_not_erase_uncovered_compute(self):
        for state in ('stopped', 'terminated', 'paused', 'stopping'):
            with self.subTest(state=state):
                r = build([resource(state=state, created_at='2026-01-01'),
                           actual('10', end='2026-09-16')],
                          estimate_hours_per_day='8', mode='combined')['resources'][0]
                cost = r['months']['2026-09']
                self.assertEqual(cost['billed'], 10)
                self.assertEqual(cost['estimated'], Decimal('14.4'))
                self.assertEqual(cost['amount'], Decimal('24.4'))

    def test_current_stopped_compute_keeps_storage_and_stopping_is_not_zeroed(self):
        model = build([resource(state='stopped', related_resource_ids=['vol-old']),
                       resource('vol-old', 'ebs', 'volume', storage_gib=100,
                                storage_type='gp2'), quote('vol-old')], mode='inventory')
        machine = next(r for r in model['resources'] if r['service'] == 'ec2')
        self.assertEqual(machine['months']['2026-09']['amount'], 0)
        self.assertEqual(machine['plan']['runtime_hours_per_day'], 0)
        self.assertEqual(machine['totals']['2026-09']['amounts']['USD'], 8)
        stopping = build([resource(state='stopping')], mode='inventory')['resources'][0]
        self.assertEqual(stopping['months']['2026-09']['amount'], Decimal('86.4'))

    def test_continuously_allocated_nodes_ignore_compute_schedule(self):
        for service, fields in [('elasticache', {'node_count': 2}),
                                ('opensearch', {'cluster_config': {'InstanceCount': 2}})]:
            with self.subTest(service=service):
                r = build([resource('cluster', service, 'cluster', details=fields)],
                          estimate_hours_per_day='8')['resources'][0]
                self.assertEqual(r['months']['2026-09']['amount'], Decimal('172.8'))
                self.assertEqual(r['plan']['runtime_hours_per_day'], 24)
                self.assertIn('with 24 runtime hours/day', ' '.join(r['plan']['notes']))

    def test_inventory_window_starts_exactly_at_midnight(self):
        window = month_windows({}, '2026-09-20T12:10:11.987654Z')[0]
        self.assertEqual(window['start'], '2026-09-01T00:00:00+00:00')
        self.assertEqual(window['hours'], 720)

    def test_parent_born_midmonth_gets_only_half_of_existing_volume_cost(self):
        model = build([resource(created_at='2026-09-16', related_resource_ids=['vol-old']),
                       resource('vol-old', 'ebs', 'volume', created_at='2026-01-01',
                                storage_gib=100, storage_type='gp2'), quote('vol-old')])
        machine = next(r for r in model['resources'] if r['service'] == 'ec2')
        volume = next(r for r in model['resources'] if r['service'] == 'ebs')
        self.assertEqual(machine['months']['2026-09']['amount'], Decimal('43.2'))
        total = machine['totals']['2026-09']
        self.assertEqual(total['components'][volume['key']]['amount'], 4)
        self.assertEqual(total['amounts']['USD'], Decimal('47.2'))
        self.assertEqual(volume['months']['2026-09']['amount'], 8)
        self.assertEqual(model['summary']['2026-09']['USD']['total'], Decimal('51.2'))

    def test_disk_attachment_date_bounds_parent_but_not_volume_own_cost(self):
        for source in ('volume', 'instance'):
            with self.subTest(source=source):
                machine = resource(created_at='2026-01-01', related_resource_ids=['vol-old'])
                volume = resource('vol-old', 'ebs', 'volume', created_at='2026-01-01',
                                  storage_gib=100, storage_type='gp2', attached_to=['i-one'])
                if source == 'volume':
                    volume['details'] = {'attachments': [{'InstanceId': 'i-one', 'AttachTime': '2026-09-16'}]}
                else:
                    machine['details'] = {'block_devices': [{'volume_id': 'vol-old', 'attached_at': '2026-09-16'}]}
                model = build([machine, volume, quote('vol-old')], period_start='2026-08-01')
                machine = next(r for r in model['resources'] if r['service'] == 'ec2')
                volume = next(r for r in model['resources'] if r['service'] == 'ebs')
                self.assertNotIn(volume['key'], machine['totals']['2026-08']['included_keys'])
                self.assertEqual(machine['totals']['2026-09']['components'][volume['key']]['amount'], 4)
                self.assertEqual(volume['months']['2026-09']['amount'], 8)

    def test_cluster_dependency_cannot_bypass_worker_or_attachment_dates(self):
        model = build([
            resource('cluster', 'eks', 'cluster', created_at='2026-01-01'),
            resource(created_at='2026-09-11', cluster_name=['cluster']),
            resource('vol-old', 'ebs', 'volume', created_at='2026-01-01',
                     cluster_name=['cluster'], attached_to=['i-one'], storage_gib=100,
                     storage_type='gp2', details={'attachments': [
                         {'InstanceId': 'i-one', 'AttachTime': '2026-09-16'}]}), quote('vol-old'),
        ])
        cluster = next(r for r in model['resources'] if r['service'] == 'eks')
        volume = next(r for r in model['resources'] if r['service'] == 'ebs')
        self.assertNotIn(volume['key'], cluster['children'])
        self.assertEqual(cluster['totals']['2026-09']['components'][volume['key']]['amount'], 4)

    def test_shared_component_workbook_prefers_union_of_parent_intervals(self):
        model = build([
            resource('i-a', created_at='2026-09-16', related_resource_ids=['vol-shared']),
            resource('i-b', created_at='2026-09-01', related_resource_ids=['vol-shared']),
            resource('vol-shared', 'ebs', 'volume', created_at='2026-01-01',
                     storage_gib=100, storage_type='gp2'), quote('vol-shared'),
        ])
        volume = next(r for r in model['resources'] if r['service'] == 'ebs')
        for style in ('readable', 'basic'):
            book = load_workbook(BytesIO(build_service_workbook(model, 'ec2', style=style)))
            rows = [row for row in book['_members'] if row[0].value == volume['key']]
            self.assertEqual([row[2].value for row in rows], [8, 4])
            self.assertIn('A5', rows[0][1].value)
            self.assertIn('A4', rows[1][1].value)

    def test_billed_interval_crossing_parent_birth_is_not_arbitrarily_split(self):
        for bill, expected in [(actual('20'), None),
                               (actual('20', start='2026-09-16'), Decimal('20'))]:
            with self.subTest(start=bill['period_start']):
                model = build([
                    resource('domain', 'emr', 'cluster', created_at='2026-09-16',
                             related_resource_ids=['i-one']),
                    resource(created_at='2026-01-01'), bill,
                ])
                parent = next(r for r in model['resources'] if r['service'] == 'emr')
                child = next(r for r in model['resources'] if r['service'] == 'ec2')
                contribution = parent['totals']['2026-09']['components'][child['key']]
                self.assertEqual(contribution['amount'], expected)
                self.assertEqual(child['months']['2026-09']['billed'], 20)
                if expected is None:
                    self.assertFalse(parent['totals']['2026-09']['complete'])

    def test_billing_covering_entire_known_lifetime_is_complete_without_catalog(self):
        r = build([resource(created_at='2026-09-16', list_price_per_hour=None),
                   actual('20', start='2026-09-16')])['resources'][0]
        cost = r['months']['2026-09']
        self.assertEqual(cost['amount'], 20)
        self.assertEqual(cost['basis'], 'billed')
        self.assertEqual(cost['modeled_hours'], 0)
        self.assertTrue(cost['complete'])

    def test_creation_after_period_has_zero_own_and_no_attached_cost(self):
        model = build([
            resource(created_at='2026-10-05T00:00:00Z', related_resource_ids=['vol-old']),
            resource('vol-old', 'ebs', 'volume', created_at='2026-01-01', storage_gib=100,
                     storage_type='gp2'), quote('vol-old'),
        ])
        machine = next(r for r in model['resources'] if r['service'] == 'ec2')
        volume = next(r for r in model['resources'] if r['service'] == 'ebs')
        entry = machine['months']['2026-09']
        self.assertEqual(entry['amount'], 0)
        self.assertEqual(entry['modeled_hours'], 0)
        self.assertTrue(entry['not_created'])
        self.assertEqual(machine['totals']['2026-09']['amounts']['USD'], 0)
        self.assertNotIn(volume['key'], machine['totals']['2026-09']['included_keys'])
        self.assertEqual(volume['months']['2026-09']['amount'], 8)
        for style in ('readable', 'basic'):
            book = load_workbook(BytesIO(build_service_workbook(model, 'ec2', style=style)))
            self.assertEqual(book['By resource']['F4'].value, 0)
            self.assertEqual(book['Attached costs']['E4'].value, 0)
            excluded = next(row for row in book['_members'] if row[0].value == volume['key'])
            self.assertEqual(excluded[1].value, '=0')

    def test_partial_month_uses_creation_in_utc_and_full_calendar_denominator(self):
        model = build([
            resource(created_at='2026-09-16T02:00:00+02:00'),
            resource('vol-new', 'ebs', 'volume', created_at='2026-09-16T00:00:00Z',
                     storage_gib=100, storage_type='gp2'), quote('vol-new'),
            resource('db', 'rds', 'db_instance', created_at='2026-09-16'),
        ])
        for r in model['resources']:
            entry = r['months']['2026-09']
            self.assertEqual(entry['modeled_hours'], 360)
            self.assertEqual(entry['excluded_hours'], 360)
            self.assertEqual(entry['amount'], Decimal('4') if r['service'] == 'ebs' else Decimal('43.2'))

    def test_shared_volume_is_counted_only_for_eligible_parent_months(self):
        model = build([
            resource('i-a', created_at='2026-09-01', related_resource_ids=['vol-shared']),
            resource('i-b', created_at='2026-10-05', related_resource_ids=['vol-shared']),
            resource('vol-shared', 'ebs', 'volume', created_at='2026-01-01',
                     storage_gib=100, storage_type='gp2'), quote('vol-shared'),
        ], period_start='2026-08-01')
        book = load_workbook(BytesIO(build_service_workbook(model, 'ec2')))
        machine = next(r for r in model['resources'] if r['id'] == 'i-a')
        self.assertEqual(machine['totals']['2026-08']['amounts']['USD'], 0)
        self.assertEqual(machine['totals']['2026-09']['amounts']['USD'], Decimal('94.4'))
        volume = next(r for r in model['resources'] if r['service'] == 'ebs')
        members = [row for row in book['_members'] if row[0].value == volume['key']]
        self.assertTrue(all(row[1].value == '=0' for row in members if row[3].value == '2026-08'))
        self.assertEqual(sum(row[1].value.startswith('=SUBTOTAL') for row in members if row[3].value == '2026-09'), 1)
        unique = next(row for row in book['_unique'] if row[0].value == volume['key'])
        self.assertIn('INDEX(_members!$C$', unique[2].value)

    def test_exclusive_end_and_month_boundaries(self):
        r = build([resource(created_at='2026-09-01T00:00:00Z')],
                  period_start='2026-08-01')['resources'][0]
        self.assertEqual(r['months']['2026-08']['amount'], 0)
        self.assertEqual(r['months']['2026-09']['amount'], Decimal('86.4'))

    def test_billing_crossing_creation_is_not_prorated_or_double_counted(self):
        r = build([resource(created_at='2026-09-16T12:00:00Z'),
                   actual('5', start='2026-09-16', end='2026-09-17')])['resources'][0]
        cost = r['months']['2026-09']
        self.assertEqual(cost['billed'], 5)
        self.assertEqual(cost['covered_hours'], 24)
        self.assertEqual(cost['modeled_hours'], 336)
        self.assertEqual(cost['estimated'], Decimal('40.32'))
        self.assertEqual(cost['amount'], Decimal('45.32'))
        self.assertEqual(cost['basis'], 'mixed')

    def test_billing_before_inventory_creation_is_preserved(self):
        r = build([resource(created_at='2026-10-02'), actual('19')])['resources'][0]
        self.assertEqual(r['months']['2026-09']['amount'], 19)
        self.assertEqual(r['months']['2026-09']['basis'], 'billed')
        self.assertFalse(r['months']['2026-09']['not_created'])

    def test_last_launch_is_not_treated_as_proof_of_nonexistence(self):
        r = build([resource(last_launched_at='2026-10-02')])['resources'][0]
        cost = r['months']['2026-09']
        self.assertIsNone(cost['amount'])
        self.assertFalse(cost['complete'])
        self.assertFalse(cost['not_created'])
        self.assertFalse(cost['allow_attachments'])
        self.assertFalse(r['lifetime']['creation_known'])
        initial = build([resource(created_at='2026-07-01', last_launched_at='2026-10-02')])['resources'][0]
        self.assertEqual(initial['months']['2026-09']['amount'], Decimal('86.4'))

    def test_missing_or_invalid_dates_are_explicit_unknown_scenarios(self):
        for created in (None, 'invalid-date'):
            r = build([resource(created_at=created)])['resources'][0]
            self.assertEqual(r['lifetime'], {'start': None, 'creation_known': False})
            self.assertEqual(r['months']['2026-09']['amount'], Decimal('86.4'))

    def test_collector_selects_primary_interface_not_list_order_or_latest_launch(self):
        ctx = _context({'ec2'})
        instance = {'InstanceId': 'i-test', 'LaunchTime': '2026-10-02',
                    'BlockDeviceMappings': [{'Ebs': {'VolumeId': 'vol-test', 'AttachTime': '2026-09-16'}}],
                    'NetworkInterfaces': [
                        {'Attachment': {'DeviceIndex': 1, 'AttachTime': '2026-10-01'}},
                        {'Attachment': {'DeviceIndex': 0, 'AttachTime': '2026-07-01'}},
                    ]}
        def pages(service, operation, region, result_key, **kwargs):
            return [{'Instances': [instance]}] if operation == 'describe_instances' else []
        with patch.object(ctx, 'safe_pages', side_effect=pages):
            collect_ec2(ctx, 'eu-west-1')
        self.assertEqual(ctx.rows[0]['created_at'], '2026-07-01')
        self.assertEqual(ctx.rows[0]['last_launched_at'], '2026-10-02')
        self.assertEqual(ctx.rows[0]['details']['block_devices'][0]['attached_at'], '2026-09-16')
