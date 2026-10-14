"""Deleted NAT gateways cannot generate estimates after their recorded deletion."""
from decimal import Decimal
from io import BytesIO
from unittest import TestCase
from unittest.mock import patch

from openpyxl import load_workbook

from aws_inventory.estimates import monthly_cost
from aws_inventory.inventory_compute import collect_vpc
from aws_inventory.service_exports import build_service_workbook
from tests.test_collectors import _context
from tests.test_html_report import actual, build, quote, resource


def nat(**fields):
    return resource('nat-old', 'vpc', 'nat_gateway', **{
        'state': 'deleted', 'created_at': '2026-01-01',
        'details': {'DeleteTime': '2026-09-16T00:00:00Z'}, **fields})


def nat_quote():
    return quote('nat-old', '0.045', 'Hrs', 'vpc')


class NatLifetimeTests(TestCase):
    def test_sdk_deletion_date_reaches_cost_model(self):
        ctx = _context({'vpc'})
        def pages(service, operation, region, result_key, **kwargs):
            if operation == 'describe_nat_gateways':
                return [{'NatGatewayId': 'nat-old', 'State': 'deleted',
                         'CreateTime': '2026-01-01', 'DeleteTime': '2026-09-16T00:00:00Z'}]
            return []
        with patch.object(ctx, 'safe_pages', side_effect=pages):
            collect_vpc(ctx, 'eu-west-1')
        self.assertEqual(ctx.rows[0]['details']['DeleteTime'], '2026-09-16T00:00:00Z')
        r = build([*ctx.rows, nat_quote()])['resources'][0]
        self.assertEqual(r['months']['2026-09']['amount'], Decimal('16.2'))
        self.assertEqual(r['months']['2026-09']['modeled_hours'], 360)
        self.assertEqual(r['months']['2026-09']['excluded_hours'], 360)

    def test_creation_deletion_and_requested_window_are_intersected(self):
        r = build([nat(created_at='2026-09-11T02:00:00+02:00'), nat_quote()],
                  period_start='2026-09-10', period_end='2026-09-20')['resources'][0]
        self.assertEqual(r['months']['2026-09']['amount'], Decimal('5.4'))
        self.assertEqual(r['months']['2026-09']['modeled_hours'], 120)
        self.assertEqual(r['months']['2026-09']['excluded_hours'], 120)

    def test_after_deletion_zero_does_not_require_a_catalog_rate(self):
        for rates in ([], [nat_quote()]):
            model = build([nat(), *rates], period_start='2026-10-01', period_end='2026-11-01')
            r = model['resources'][0]
            cost = r['months']['2026-10']
            self.assertEqual(cost['amount'], 0)
            self.assertTrue(cost['after_deletion'])
            self.assertFalse(cost['not_created'])
            self.assertFalse(cost['allow_attachments'])
            for style in ('readable', 'basic'):
                book = load_workbook(BytesIO(build_service_workbook(model, 'vpc', style=style)))
                self.assertEqual(book['By resource']['F4'].value, 0)
                self.assertIn('already deleted', book['By resource']['E4'].value)
                self.assertIn('already deleted', book['Attached costs']['D4'].value)

    def test_missing_invalid_or_inverted_deletion_dates_are_unavailable(self):
        for value in (None, 'invalid', '2025-12-31'):
            r = build([nat(details={'DeleteTime': value}), nat_quote()])['resources'][0]
            self.assertIsNone(r['months']['2026-09']['amount'])
            self.assertNotIn('end', r['lifetime'])

    def test_deleting_state_does_not_prove_deletion(self):
        r = build([nat(state='deleting'), nat_quote()])['resources'][0]
        self.assertNotIn('end', r['lifetime'])
        self.assertEqual(r['months']['2026-09']['amount'], Decimal('32.4'))

    def test_parent_subtotal_and_excel_use_the_clipped_component(self):
        model = build([nat(parent_id='vpc-one'), nat_quote(),
                       resource('vpc-one', 'vpc', 'vpc', list_price_per_hour=None)])
        gateway = next(r for r in model['resources'] if r['id'] == 'nat-old')
        parent = next(r for r in model['resources'] if r['id'] == 'vpc-one')
        self.assertEqual(parent['totals']['2026-09']['components'][gateway['key']]['amount'], Decimal('16.2'))
        self.assertEqual(model['summary']['2026-09']['USD']['total'], Decimal('16.2'))
        book = load_workbook(BytesIO(build_service_workbook(model, 'vpc')))
        members = [row for row in book['_members'] if row[0].value == gateway['key']]
        self.assertEqual([row[2].value for row in members], [16.2, 16.2])

    def test_billing_crossing_deletion_is_preserved_without_estimate_overlap(self):
        model = build([nat(), nat_quote()])
        r = model['resources'][0]
        cost = monthly_cost(r['plan'], model['windows'][0],
                            [actual('5', start='2026-09-15', end='2026-09-17')], r['lifetime'])
        self.assertEqual(cost['billed'], 5)
        self.assertEqual(cost['modeled_hours'], 336)
        self.assertEqual(cost['amount'], Decimal('20.12'))
