"""End-to-end checks of offline per-service workbooks and companion CSV tables."""

from __future__ import annotations

import csv
import json
import tempfile
import unittest
from collections import Counter
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from examples.simulate_export import (
    ACCOUNT, BUNDLE_START, CENTRAL, CLUSTER, END, NODE_A, NODE_B, STANDALONE,
    UNATTACHED, WEST, run_simulation,
)


MONTHS = ("2026-07", "2026-08", "2026-09")


def amount(value):
    return Decimal(str(value).replace(",", "."))


class ServiceBundleSimulationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.directory = tempfile.TemporaryDirectory(prefix="costcsv-service-bundle-")
        cls.addClassCleanup(cls.directory.cleanup)
        cls.output = Path(cls.directory.name) / "reports"
        cls.result = run_simulation(cls.output, report_format="bundle")
        if cls.result["exit_code"] != 0:
            raise AssertionError(cls.result["stderr"])
        cls.manifest = json.loads((cls.output / "manifest.json").read_text(encoding="utf-8"))
        cls.reports = {report["service"]: report for report in cls.manifest["reports"]}

    def table(self, service, name):
        entry = next(table for table in self.reports[service]["tables"] if table["name"] == name)
        with (self.output / entry["csv"]).open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream, delimiter=";")
            return reader.fieldnames, list(reader)

    def test_service_files_and_csvs_are_separate_and_explicitly_simulated(self):
        self.assertEqual(self.result["network_attempts"], 0)
        self.assertTrue(self.manifest["simulation_only"])
        self.assertEqual(self.manifest["status"], "complete")
        self.assertEqual(set(self.reports), {"overview", "ec2", "ec2_other", "ebs", "eks"})
        self.assertEqual({report["workbook"] for report in self.reports.values()}, {
            "00_overview.xlsx", "ec2.xlsx", "ec2_other.xlsx", "ebs.xlsx", "eks.xlsx",
        })
        expected_files = set(self.manifest["files"]) | {"manifest.json"}
        actual_files = {path.relative_to(self.output).as_posix()
                        for path in self.output.rglob("*") if path.is_file()}
        self.assertEqual(actual_files, expected_files)
        for service, report in self.reports.items():
            for table in report["tables"]:
                path = self.output / table["csv"]
                self.assertEqual(path.parent, self.output / "csv" / service)
                self.assertTrue(path.read_bytes().startswith(b"\xef\xbb\xbf"))
                headers, rows = self.table(service, table["name"])
                self.assertEqual(headers[0], "Simulation only")
                self.assertTrue(all(row["Simulation only"] == "true" for row in rows))
                self.assertEqual(len(rows), table["rows"])

    def test_three_month_totals_match_primary_costs_without_duplicate_details(self):
        headers, rows = self.table("overview", "Monthly total")
        self.assertTrue(all(month in headers for month in MONTHS))
        self.assertEqual(len(rows), 1)
        total = rows[0]
        self.assertEqual(total["Currency"], "USD")
        self.assertEqual(total["Cost metric"], "AmortizedCost")
        self.assertEqual([amount(total[month]) for month in MONTHS],
                         [Decimal("261.12"), Decimal("273.62"), Decimal("274.72")])
        self.assertEqual(amount(total["Period total"]), Decimal("809.46"))
        primary = [row for row in self.result["rows"] if row["row_type"] == "cost_summary"]
        details = [row for row in self.result["rows"] if row["row_type"] == "cost_breakdown"]
        self.assertEqual(len(primary), 15)
        self.assertEqual(len(details), 12)
        self.assertEqual(sum(Decimal(row["amount"]) for row in primary), Decimal("809.46"))
        self.assertEqual(Counter(row["cost_scope"] for row in details),
                         {"ec2_instance_type": 6, "ebs_region": 6})
        service_total = Decimal(0)
        for service in ("ec2", "ec2_other", "eks"):
            _, service_rows = self.table(service, "Monthly by region")
            service_total += sum(amount(row["Period total"]) for row in service_rows)
        self.assertEqual(service_total, Decimal("809.46"))
        _, ebs = self.table("ebs", "Monthly by region")
        self.assertEqual(sum(amount(row["Period total"]) for row in ebs), Decimal("36.86"))
        central = next(row for row in ebs if row["Region"] == CENTRAL)
        self.assertEqual(central["2026-09"], "-1,20")

    def test_ec2_type_breakdown_is_exact_and_resource_costs_remain_unallocated(self):
        _, types = self.table("ec2", "Monthly by instance type")
        self.assertEqual(len(types), 2)
        by_type = {row["Instance type"]: row for row in types}
        self.assertEqual(amount(by_type["m6i.large"]["Period total"]), Decimal("502.80"))
        self.assertEqual(amount(by_type["t3.small"]["Period total"]), Decimal("49.00"))
        self.assertEqual(by_type["m6i.large"]["Region"], WEST)
        self.assertEqual(by_type["t3.small"]["Region"], CENTRAL)
        _, resources = self.table("ec2", "Monthly by instance")
        self.assertEqual({row["Resource ID"] for row in resources}, {NODE_A, NODE_B, STANDALONE})
        for row in resources:
            self.assertTrue(all(row[month] == "" for month in MONTHS))
            self.assertEqual(row["Period total"], "")
            self.assertIn("Unavailable", row["Data status"])
        _, clusters = self.table("eks", "Monthly by cluster")
        self.assertEqual(len(clusters), 1)
        self.assertEqual(clusters[0]["Name"], CLUSTER)
        self.assertTrue(all(clusters[0][month] == "" for month in MONTHS))

    def test_inventory_names_cluster_links_and_disks_survive_readable_views(self):
        _, instances = self.table("ec2", "Instances")
        by_id = {row["Resource ID"]: row for row in instances}
        self.assertEqual(by_id[NODE_A]["Name"], "worker-a")
        self.assertEqual(by_id[NODE_A]["Storage (GiB)"], "132")
        self.assertEqual(by_id[NODE_B]["Name"], NODE_B)
        self.assertEqual(by_id[NODE_B]["Storage (GiB)"], "32")
        self.assertEqual(by_id[STANDALONE]["Name"], "batch-runner")
        _, nodes = self.table("eks", "Cluster nodes")
        self.assertEqual({row["Resource ID"] for row in nodes}, {NODE_A, NODE_B})
        self.assertTrue(all(row["Cluster"] == CLUSTER for row in nodes))
        _, disks = self.table("ec2", "Attached disks")
        self.assertEqual(len(disks), 4)
        self.assertEqual(sum(amount(row["Storage (GiB)"]) for row in disks), Decimal(184))
        _, unattached = self.table("ebs", "Unattached volumes")
        self.assertEqual(len(unattached), 1)
        self.assertEqual(unattached[0]["Resource ID"], UNATTACHED)
        self.assertEqual(unattached[0]["Storage (GiB)"], "200")

    def test_excel_money_is_numeric_and_matches_csv_without_formula_cells(self):
        for service, report in self.reports.items():
            workbook = load_workbook(self.output / report["workbook"], read_only=True, data_only=False)
            try:
                self.assertEqual(workbook.sheetnames, [table["name"] for table in report["tables"]])
                for sheet in workbook:
                    self.assertIn("SIMULATED DATA", sheet.cell(2, 1).value)
                    headers = [cell.value for cell in sheet[5]]
                    self.assertNotIn("row_type", headers)
                    for row in sheet.iter_rows():
                        self.assertTrue(all(cell.data_type != "f" for cell in row))
                    if all(month in headers for month in MONTHS):
                        _, companion = self.table(service, sheet.title)
                        excel_rows = list(sheet.iter_rows(min_row=6))
                        self.assertEqual(len(excel_rows), len(companion))
                        for cells, csv_row in zip(excel_rows, companion):
                            for title in (*MONTHS, "Period total"):
                                value = cells[headers.index(title)]
                                if csv_row[title] == "":
                                    self.assertIsNone(value.value)
                                else:
                                    self.assertEqual(value.data_type, "n")
                                    self.assertEqual(amount(value.value), amount(csv_row[title]))
            finally:
                workbook.close()

    def test_breakdown_queries_use_the_entire_period_and_account_boundary(self):
        self.assertEqual(len(self.result["calls"]), 30)
        costs = [call for call in self.result["calls"] if call["service"] == "ce"]
        self.assertEqual(len(costs), 5)
        for call in costs:
            request = call["parameters"]
            self.assertEqual(request["TimePeriod"], {"Start": BUNDLE_START, "End": END})
            encoded_filter = json.dumps(request["Filter"])
            self.assertIn("LINKED_ACCOUNT", encoded_filter)
            self.assertIn(ACCOUNT, encoded_filter)
        instance_types = costs[2]["parameters"]
        self.assertEqual([group["Key"] for group in instance_types["GroupBy"]],
                         ["REGION", "INSTANCE_TYPE"])
        self.assertEqual(costs[3]["operation"], "GetDimensionValues")
        self.assertEqual(costs[3]["parameters"]["Dimension"], "USAGE_TYPE")
        self.assertEqual([group["Key"] for group in costs[4]["parameters"]["GroupBy"]],
                         ["REGION", "SERVICE"])


class PartialBundleSimulationTests(unittest.TestCase):
    def test_denied_inventory_read_produces_partial_bundle_or_no_strict_output(self):
        with tempfile.TemporaryDirectory(prefix="costcsv-partial-bundle-") as directory:
            destination = Path(directory) / "partial"
            result = run_simulation(destination, report_format="bundle", deny_snapshots=True)
            self.assertEqual(result["exit_code"], 2, result["stderr"])
            manifest = json.loads((destination / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["status"], "partial")
            self.assertTrue(manifest["simulation_only"])
            error_rows = [row for row in result["rows"] if row["row_type"] == "error"]
            self.assertEqual(len(error_rows), 1)
            self.assertEqual(error_rows[0]["error_code"], "UnauthorizedOperation")
            strict_destination = Path(directory) / "strict"
            strict_result = run_simulation(strict_destination, report_format="bundle",
                                           deny_snapshots=True, strict=True)
            self.assertEqual(strict_result["exit_code"], 1, strict_result["stderr"])
            self.assertFalse(strict_destination.exists())
            self.assertEqual(sorted(path.name for path in Path(directory).iterdir()), ["partial"])


if __name__ == "__main__":
    unittest.main()
