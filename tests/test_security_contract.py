"""Security boundaries: metadata projection, safe reports and AWS read allowlists."""
from copy import deepcopy
from io import BytesIO
import base64
import hashlib
import json
from pathlib import Path
import re
import unittest
from unittest.mock import Mock, patch
from zipfile import ZipFile

import boto3
from botocore.stub import Stubber
from openpyxl import load_workbook

from aws_inventory.common import InventoryContext, compact_json
from aws_inventory.inventory_services import collect_lambda, collect_cloudfront
from aws_inventory.inventory_extended import collect_glue
from aws_inventory.html_report import render_html, SCRIPT
from aws_inventory.read_operations import OPERATIONS, example_policy
from aws_inventory.report import redact_sensitive
from aws_inventory.service_exports import build_service_workbook
from examples.simulate_export import offline_environment, ACCOUNT, GENERATED, WEST
from tests.test_html_report import build, resource

ROOT = Path(__file__).resolve().parents[1]


class SecurityContractTests(unittest.TestCase):
    def context(self, session=None):
        return InventoryContext(session, ACCOUNT, 'aws', GENERATED, {'lambda'}, {WEST}, False)

    def assert_not_exported(self, rows, markers):
        serialized = compact_json(rows)
        model = build(rows)
        document = render_html(model)
        payload = json.loads(re.search(r'<script id="report-data" type="application/json">(.*?)</script>', document, re.S)[1])
        for marker in markers:
            self.assertNotIn(marker, serialized)
            self.assertNotIn(marker, document)
        for export in payload['workbooks'].values():
            with ZipFile(BytesIO(base64.b64decode(export['base64']))) as archive:
                for name in archive.namelist():
                    content = archive.read(name).decode('utf-8')
                    for marker in markers:
                        self.assertNotIn(marker, content)

    def test_lambda_environment_received_in_memory_is_not_exported(self):
        arn = f'arn:aws:lambda:{WEST}:{ACCOUNT}:function:worker'
        function = {'FunctionName': 'worker', 'FunctionArn': arn, 'Runtime': 'python3.11',
                    'Environment': {'Variables': {'ORDINARY_NAME': 'PRIVATE-ENV-CANARY'}},
                    'MemorySize': 128, 'Timeout': 3}
        original = deepcopy(function)
        with offline_environment() as attempts:
            for fallback in (False, True):
                with self.subTest(fallback=fallback):
                    ctx = self.context(boto3.Session())
                    client = ctx.client('lambda', WEST)
                    with Stubber(client) as stub:
                        stub.add_response('list_functions', {'Functions': [function]}, {})
                        stub.add_response('list_tags', {'Tags': {}}, {'Resource': arn})
                        if fallback:
                            stub.add_client_error('get_function_configuration', 'AccessDeniedException',
                                                  expected_params={'FunctionName': 'worker'})
                        else:
                            stub.add_response('get_function_configuration', function, {'FunctionName': 'worker'})
                        collect_lambda(ctx, WEST)
                        stub.assert_no_pending_responses()
                    self.assertEqual(ctx.rows[0]['memory_mib'], 128)
                    self.assert_not_exported(ctx.rows, ['PRIVATE-ENV-CANARY', 'ORDINARY_NAME'])
            self.assertEqual(attempts, [])
        self.assertEqual(function, original)

    def test_cloudfront_headers_and_glue_arguments_are_projected_out(self):
        distribution = {'Id': 'DIST1', 'DomainName': 'distribution.example', 'Origins': {'Items': [
            {'Id': 'origin', 'DomainName': 'origin.example',
             'CustomHeaders': {'Items': [{'HeaderName': 'X-Key', 'HeaderValue': 'PRIVATE-HEADER-CANARY'}]}}]}}
        job = {'Name': 'job', 'WorkerType': 'G.1X', 'NumberOfWorkers': 2,
               'DefaultArguments': {'--password': 'PRIVATE-JOB-CANARY'},
               'NonOverridableArguments': {'--token': 'PRIVATE-ARG-CANARY'}}
        with offline_environment() as attempts:
            ctx = self.context()
            with patch.object(ctx, 'safe_pages', return_value=[distribution]):
                collect_cloudfront(ctx, WEST)
            with patch.object(ctx, 'safe_pages', return_value=[job]):
                collect_glue(ctx, WEST)
            self.assert_not_exported(ctx.rows, ['PRIVATE-HEADER-CANARY', 'PRIVATE-JOB-CANARY', 'PRIVATE-ARG-CANARY'])
            self.assertEqual(attempts, [])

    def test_unknown_operations_and_secret_reads_fail_before_client_creation(self):
        session = Mock()
        ctx = self.context(session)
        blocked = [('ec2', 'terminate_instances'), ('s3', 'put_object'), ('s3', 'get_object'),
                   ('s3', 'list_objects_v2'), ('secretsmanager', 'get_secret_value'),
                   ('sts', 'get_session_token'), ('sts', 'get_federation_token'),
                   ('ec2', 'describe_instance_attribute'), ('unknown', 'list_items')]
        for service, operation in blocked:
            for method in ('call', 'pages', 'safe_call'):
                with self.subTest(service=service, operation=operation, method=method):
                    with self.assertRaisesRegex(ValueError, 'allowlist'):
                        if method == 'pages':
                            list(ctx.pages(service, operation, WEST, 'Items'))
                        else:
                            getattr(ctx, method)(service, operation, WEST)
        session.client.assert_not_called()

    def test_cached_sdk_client_also_blocks_unapproved_operations(self):
        with offline_environment() as attempts:
            ctx = self.context(boto3.Session())
            for service, operation, parameters in [
                ('ec2', 'delete_volume', {'VolumeId': 'vol-1234567890abcdef0'}),
                ('secretsmanager', 'get_secret_value', {'SecretId': 'private'}),
                ('sts', 'get_session_token', {}),
            ]:
                client = ctx.client(service, WEST)
                with self.subTest(operation=operation), patch.object(client._endpoint, 'make_request') as request:
                    with self.assertRaisesRegex(ValueError, 'allowlist'):
                        getattr(client, operation)(**parameters)
                    request.assert_not_called()
            self.assertEqual(attempts, [])

    def test_policy_matches_allowlist_without_wildcard_actions(self):
        policy = json.loads((ROOT / 'iam/inventory-read-only.json').read_text())
        self.assertEqual(policy, example_policy())
        actions = {action for statement in policy['Statement'] for action in statement['Action']}
        self.assertEqual(actions, {action for operations in OPERATIONS.values() for action in operations.values()})
        self.assertTrue(all('*' not in action and '?' not in action for action in actions))
        self.assertLessEqual(len(json.dumps(policy, separators=(',', ':'))), 6144)
        self.assertEqual(OPERATIONS['s3']['get_bucket_encryption'], 's3:GetEncryptionConfiguration')
        self.assertEqual(OPERATIONS['apigateway']['get_stages'], 'apigateway:GET')
        self.assertNotIn('secretsmanager:GetSecretValue', actions)

    def test_html_payload_is_inert_and_script_hash_matches(self):
        name = '</script><img src=x onerror="alert(1)"><script>alert(2)</script>'
        document = render_html(build([resource(resource_name=name, tags={'api_key': 'PRIVATE-TAG-CANARY'})]))
        self.assertNotIn(name, document)
        self.assertNotIn('PRIVATE-TAG-CANARY', document)
        self.assertEqual(document.count('<script'), 2)
        self.assertEqual(document.count('</script>'), 2)
        payload = json.loads(re.search(r'<script id="report-data" type="application/json">(.*?)</script>', document, re.S)[1])
        self.assertEqual(payload['resources'][0]['name'], name)
        digest = base64.b64encode(hashlib.sha256(SCRIPT.encode()).digest()).decode()
        self.assertIn('sha256-' + digest, document)
        self.assertNotIn('innerHTML', SCRIPT)
        self.assertNotIn('eval(', SCRIPT)

    def test_excel_user_values_remain_literal_and_have_no_external_links(self):
        names = ['=HYPERLINK("https://invalid.example","x")', '+1+1', '-1+1', '@SUM(A1:A2)',
                 "=cmd|' /C calc'!A0", 'https://invalid.example', '\t=1+1']
        model = build([resource(f'i-{index}', resource_name=name) for index, name in enumerate(names)])
        content = build_service_workbook(model, 'ec2')
        book = load_workbook(BytesIO(content))
        cells = [row[0] for row in book['By resource'].iter_rows(min_row=4) if row[0].value in names]
        self.assertEqual(len(cells), len(names))
        for cell in cells:
            self.assertEqual(cell.data_type, 's')
            self.assertIsNone(cell.hyperlink)
        with ZipFile(BytesIO(content)) as archive:
            self.assertFalse(any('externalLinks/' in name for name in archive.namelist()))

    def test_redaction_is_key_based_not_a_free_text_secret_detector(self):
        source = {'password': 'PRIVATE-PASSWORD-CANARY', 'description': 'arbitrary-free-text-canary'}
        redacted = redact_sensitive(source)
        self.assertEqual(redacted['password'], '[REDACTED]')
        self.assertEqual(redacted['description'], source['description'])
        self.assertEqual(source['password'], 'PRIVATE-PASSWORD-CANARY')
