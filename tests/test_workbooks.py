"""Offline checks for readable spreadsheets and safe report-directory publication."""

from __future__ import annotations

import csv
import json
import stat
import sys
import tempfile
import threading
import types
import unittest
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from pathlib import Path
from unittest.mock import Mock, patch
from zipfile import ZipFile

from openpyxl import load_workbook

import costcsv
from aws_inventory import workbooks


GENERATED = "2026-10-11T12:00:00+00:00"
CONTEXT = {"caller_account_id": "123456789012", "simulation_only": True}


def fixture_reports(*, amount=Decimal("-12.34000000000000000001"), text="=2+2"):
    columns = [types.SimpleNamespace(key="name", title="Instance name", kind="text"),
               types.SimpleNamespace(key="amount", title="2026-09 (USD)", kind="money"),
               types.SimpleNamespace(key="price", title="Current price / hour", kind="rate"),
               types.SimpleNamespace(key="detail", title="Notes", kind="text")]
    table = types.SimpleNamespace(
        name="Monthly by region", columns=columns,
        rows=[{"name": text, "amount": amount, "price": "0.1234567890123456789",
               "detail": {"password": "never-export-this", "source": "fixture"}}],
        note="Monthly billed costs; prices describe the current catalog.")
    return [types.SimpleNamespace(key="ec2", title="Amazon EC2", tables=[table])]


def directory_bytes(path):
    return {item.relative_to(path).as_posix(): item.read_bytes()
            for item in path.rglob("*") if item.is_file()}


class WorkbookPublicationTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="aws-workbooks-test-")
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.output = self.root / "reports"

    def export(self, *, reports=None, overwrite=False, path=None, builder=None, excel=True):
        # Test the real serializers independently of service aggregation and AWS clients.
        module = types.ModuleType("aws_inventory.report_views")
        module.build_service_reports = builder or Mock(return_value=reports or fixture_reports())
        with patch.dict(sys.modules, {"aws_inventory.report_views": module}):
            return workbooks.write_service_reports(
                path or self.output, [], generated_at=GENERATED, context=CONTEXT,
                partial=False, excel=excel, overwrite=overwrite)

    def assert_clean_parent(self, *, expected=None):
        self.assertEqual(set(self.root.iterdir()), set(expected or []))

    def test_readable_xlsx_typed_cells_formula_defense_and_exact_csv(self):
        self.export()
        self.assertEqual(stat.S_IMODE(self.output.stat().st_mode), 0o700)
        manifest = json.loads((self.output / "manifest.json").read_text())
        self.assertEqual(manifest["format"], workbooks.BUNDLE_FORMAT)
        self.assertIs(manifest["simulation_only"], True)
        workbook_path = self.output / "ec2.xlsx"
        workbook = load_workbook(workbook_path, data_only=False)
        self.addCleanup(workbook.close)
        sheet = workbook["Monthly by region"]
        self.assertEqual(sheet["A5"].value, "Instance name")
        self.assertIn("SIMULATED DATA", sheet["A2"].value)
        self.assertEqual(sheet.freeze_panes, "C6")
        self.assertEqual(sheet.auto_filter.ref, "A5:D6")
        self.assertEqual(sheet["A6"].value, "=2+2")
        self.assertEqual(sheet["A6"].data_type, "s")
        self.assertEqual(sheet["B6"].data_type, "n")
        self.assertAlmostEqual(sheet["B6"].value, -12.34)
        self.assertEqual(sheet["C6"].data_type, "n")
        self.assertNotIn("never-export-this", sheet["D6"].value)
        self.assertIn("[REDACTED]", sheet["D6"].value)
        with ZipFile(workbook_path) as archive:
            self.assertFalse(any(name.startswith("xl/externalLinks/") for name in archive.namelist()))
            worksheet_xml = archive.read("xl/worksheets/sheet1.xml")
            self.assertNotIn(b"<f>", worksheet_xml)
        csv_path = self.output / manifest["reports"][0]["tables"][0]["csv"]
        self.assertTrue(csv_path.read_bytes().startswith(b"\xef\xbb\xbf"))
        with csv_path.open(encoding="utf-8-sig", newline="") as stream:
            row, = list(csv.DictReader(stream, delimiter=";"))
        self.assertEqual(row["Instance name"], "'=2+2")
        self.assertEqual(row["2026-09 (USD)"], "-12,34000000000000000001")
        self.assertEqual(row["Current price / hour"], "0,1234567890123456789")
        for relative in manifest["files"] + ["manifest.json"]:
            self.assertEqual(stat.S_IMODE((self.output / relative).stat().st_mode), 0o600)
        workbooks.validate_bundle_destination(self.output, overwrite=True)
        self.assert_clean_parent(expected=[self.output])

    def test_all_external_formulas_and_urls_stay_text_and_controls_are_escaped(self):
        payloads = ["=HYPERLINK(\"https://example.invalid\",\"Open\")", "+1+1", "-1+1",
                    "@SUM(A1:A2)", "\t=NOW()", "https://example.invalid", "name\x00suffix"]
        reports = fixture_reports()
        table = reports[0].tables[0]
        table.rows = [{"name": payload, "amount": "-1"} for payload in payloads]
        self.export(reports=reports)
        workbook = load_workbook(self.output / "ec2.xlsx", data_only=False)
        self.addCleanup(workbook.close)
        sheet = workbook.active
        for number, payload in enumerate(payloads, start=6):
            item = sheet.cell(number, 1)
            self.assertEqual(item.data_type, "s")
            self.assertIsNone(item.hyperlink)
            self.assertEqual(item.value, payload.replace("\x00", "\\u0000"))

    def test_invalid_numeric_input_never_publishes_a_directory(self):
        for amount in (True, {}, [], "NaN", "Infinity", "-Infinity", "=SUM(A1:A2)"):
            with self.subTest(amount=amount):
                with self.assertRaises(ValueError):
                    self.export(reports=fixture_reports(amount=amount))
                self.assert_clean_parent()

    def test_csv_failure_preserves_the_previous_complete_bundle(self):
        self.export()
        previous = directory_bytes(self.output)
        with patch.object(workbooks, "_write_csv", side_effect=OSError("Simulated full disk")):
            with self.assertRaisesRegex(OSError, "full disk"):
                self.export(overwrite=True)
        self.assertEqual(directory_bytes(self.output), previous)
        self.assert_clean_parent(expected=[self.output])

    def test_failed_final_rename_restores_previous_directory(self):
        self.export()
        previous = directory_bytes(self.output)
        real_rename = Path.rename

        def fail_staging_publish(source, destination):
            if source.name.startswith(".reports.staging-"):
                raise OSError("Simulated publication failure")
            return real_rename(source, destination)

        with patch.object(Path, "rename", fail_staging_publish):
            with self.assertRaisesRegex(OSError, "publication failure"):
                self.export(overwrite=True)
        self.assertEqual(directory_bytes(self.output), previous)
        self.assert_clean_parent(expected=[self.output])

    def test_existing_bundle_requires_explicit_overwrite(self):
        self.export()
        previous = directory_bytes(self.output)
        with self.assertRaises(FileExistsError):
            self.export()
        self.assertEqual(directory_bytes(self.output), previous)
        self.export(reports=fixture_reports(text="replacement"), overwrite=True)
        self.assertNotEqual(directory_bytes(self.output), previous)
        self.assert_clean_parent(expected=[self.output])

    def test_arbitrary_directories_and_extra_files_are_never_replaced(self):
        self.output.mkdir()
        personal = self.output / "personal-notes.txt"
        personal.write_text("Keep my notes.")
        with self.assertRaisesRegex(ValueError, "manifest"):
            self.export(overwrite=True)
        self.assertEqual(personal.read_text(), "Keep my notes.")
        personal.unlink()
        self.output.rmdir()
        self.export()
        personal.write_text("Keep my notes.")
        previous = directory_bytes(self.output)
        with self.assertRaisesRegex(ValueError, "additional files"):
            self.export(overwrite=True)
        self.assertEqual(directory_bytes(self.output), previous)
        self.assert_clean_parent(expected=[self.output])

    def test_symlink_destinations_and_symlink_entries_are_never_replaced(self):
        target = self.root / "personal"
        target.mkdir()
        (target / "notes.txt").write_text("Keep this.")
        self.output.symlink_to(target, target_is_directory=True)
        with self.assertRaisesRegex(ValueError, "symbolic link"):
            self.export(overwrite=True)
        self.assertEqual((target / "notes.txt").read_text(), "Keep this.")
        self.output.unlink()
        self.export()
        csv_path = next((self.output / "csv").rglob("*.csv"))
        csv_path.unlink()
        csv_path.symlink_to(target / "notes.txt")
        with self.assertRaisesRegex(ValueError, "symbolic link"):
            self.export(overwrite=True)
        self.assertEqual((target / "notes.txt").read_text(), "Keep this.")

    def test_invalid_manifest_structures_are_refused_cleanly(self):
        self.output.mkdir()
        path = self.output / "manifest.json"
        for manifest in ([], None, "text", {"format": workbooks.BUNDLE_FORMAT, "files": "a"}):
            with self.subTest(manifest=manifest):
                path.write_text(json.dumps(manifest))
                with self.assertRaises(ValueError):
                    self.export(overwrite=True)
                self.assertEqual(json.loads(path.read_text()), manifest)

    def test_destination_created_during_staging_survives_revalidation(self):
        real_write_csv = workbooks._write_csv

        def create_unrelated_destination(*args, **kwargs):
            real_write_csv(*args, **kwargs)
            self.output.mkdir()
            (self.output / "personal.txt").write_text("Created during export.")

        with patch.object(workbooks, "_write_csv", side_effect=create_unrelated_destination):
            with self.assertRaises(ValueError):
                self.export(overwrite=True)
        self.assertEqual((self.output / "personal.txt").read_text(), "Created during export.")
        self.assert_clean_parent(expected=[self.output])

    def test_overwrite_lock_prevents_concurrent_replacement(self):
        self.export()
        building = threading.Event()
        release = threading.Event()

        def blocking_builder(*args, **kwargs):
            building.set()
            if not release.wait(5):
                raise TimeoutError("Test did not release the workbook builder")
            return fixture_reports(text="first writer")

        with ThreadPoolExecutor(max_workers=1) as executor:
            first = executor.submit(self.export, overwrite=True, builder=blocking_builder)
            try:
                self.assertTrue(building.wait(5))
                with self.assertRaisesRegex(FileExistsError, "Another export"):
                    self.export(overwrite=True)
            finally:
                release.set()
            first.result(timeout=5)
        workbook = load_workbook(self.output / "ec2.xlsx", read_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active["A6"].value, "first writer")
        self.assert_clean_parent(expected=[self.output])


class BundleArgumentTests(unittest.TestCase):
    def test_bundle_rejects_custom_cost_shape_before_any_execution(self):
        for flags in (["--granularity", "DAILY"], ["--group-by", "INSTANCE_TYPE"]):
            with self.subTest(flags=flags), patch.object(costcsv, "execute") as execute:
                code = costcsv.main(["--from", "2026-09-01", "--to", "2026-10-01", *flags])
                self.assertEqual(code, 1)
                execute.assert_not_called()

    def test_raw_csv_keeps_custom_daily_grouping(self):
        with tempfile.TemporaryDirectory(prefix="aws-csv-arguments-") as directory:
            args = costcsv.parser().parse_args([
                "--format", "raw-csv", "--granularity", "DAILY",
                "--group-by", "INSTANCE_TYPE", "--from", "2026-09-01", "--to", "2026-10-01",
                "--output", str(Path(directory) / "technical.csv"),
            ])
            costcsv.validate(args)
            self.assertEqual(args.group_by, ["INSTANCE_TYPE"])
            self.assertEqual(args.granularity, "DAILY")


if __name__ == "__main__":
    unittest.main()
