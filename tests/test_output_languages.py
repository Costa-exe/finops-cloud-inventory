"""Localized report text, prompts and workbook amounts remain consistent."""
import ast
import base64
from decimal import Decimal
from contextlib import redirect_stdout
from io import BytesIO, StringIO
import json
from pathlib import Path
import re
from string import Formatter
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

import costcsv
from aws_inventory.localization import LANGUAGES, LABELS, labels
from aws_inventory.report_template import NOTE_TRANSLATIONS
from aws_inventory.translations import LABELS as FRENCH_SPANISH, MESSAGES, TEMPLATES
from examples.simulate_export import run_simulation


class OutputLanguageTests(unittest.TestCase):
    def test_translations_cover_labels_notes_and_parameter_placeholders(self):
        self.assertEqual(LABELS.keys(), FRENCH_SPANISH.keys())
        self.assertFalse(NOTE_TRANSLATIONS.keys() - MESSAGES.keys())
        for key, values in TEMPLATES.items():
            fields = [{field for _, field, _, _ in Formatter().parse(value) if field} for value in values]
            self.assertTrue(all(field == fields[0] for field in fields), key)
        for language in LANGUAGES:
            self.assertTrue(all(labels(language).values()))

    def test_static_messages_in_html_and_excel_have_translations(self):
        root = Path(__file__).resolve().parents[1] / 'aws_inventory'
        for name in ('report_ui.js', 'report_calculation.js', 'service_exports.py'):
            source = (root / name).read_text()
            if name.endswith('.py'):
                messages = [node.args[0].value for node in ast.walk(ast.parse(source))
                            if isinstance(node, ast.Call) and isinstance(node.func, ast.Name)
                            and node.func.id == 'tr' and isinstance(node.args[0], ast.Constant)]
            else:
                messages = re.findall(r"tr\('([^']*)'", source)
            self.assertFalse(set(messages) - MESSAGES.keys(), name)

    def test_all_four_languages_are_available_in_prompt_and_cli(self):
        for language in LANGUAGES:
            with self.subTest(language=language):
                args = costcsv.parser().parse_args(['--mode', 'inventory', '--region', 'eu-west-1',
                                                   '--service', 'ec2', '--language', language])
                costcsv.validate(args)
                with patch('sys.stdin.isatty', return_value=True), redirect_stdout(StringIO()), \
                        patch('builtins.input', side_effect=[language, 'dmy', 'Europe/Rome', '24']) as prompt:
                    costcsv.guided(args)
                self.assertEqual(args.language, language)
                self.assertIn('en/it/fr/es', prompt.call_args_list[0].args[0])

    def test_french_and_spanish_offline_exports_preserve_costs_and_names(self):
        for language, overview, own_sheet, attached_sheet, account in [
            ('fr', 'Vue d’ensemble', 'Par ressource', 'Coûts associés', 'Compte'),
            ('es', 'Resumen', 'Por recurso', 'Costes asociados', 'Cuenta'),
        ]:
            with self.subTest(language=language), tempfile.TemporaryDirectory() as directory:
                result = run_simulation(report_format='html', working_directory=directory, language=language)
                self.assertEqual(result['exit_code'], 0, result['stderr'])
                self.assertEqual(result['network_attempts'], 0)
                document = Path(result['output']).read_text()
                self.assertIn(f'<html lang="{language}">', document)
                model = json.loads(re.search(r'<script id="report-data" type="application/json">(.*?)</script>', document, re.S)[1])
                self.assertEqual(model['labels']['overview'], overview)
                self.assertEqual(Decimal(model['summary']['2026-09']['USD']['total']), Decimal('297.92'))
                workbook = load_workbook(BytesIO(base64.b64decode(model['workbooks']['ec2']['readable']['base64'])))
                self.assertIn(own_sheet, workbook.sheetnames)
                self.assertIn(attached_sheet, workbook.sheetnames)
                self.assertEqual(workbook[own_sheet]['A2'].value, f'{account} 123456789012')
                self.assertTrue(all(len(sheet.title) <= 31 for sheet in workbook))
                worker = next(row for row in workbook[own_sheet].iter_rows(min_row=4)
                              if row[0].value == 'worker-a')
                # July predates the only known EC2 launch; September remains fully estimated.
                self.assertIsNone(worker[5].value)
                self.assertEqual(worker[7].value, 86.4)
