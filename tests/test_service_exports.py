"""Monthly Excel downloads, shared costs and descriptive resource links."""
from decimal import Decimal
import base64
from io import BytesIO
import unittest

from openpyxl import load_workbook

from aws_inventory.service_exports import build_service_workbook, embedded_workbooks
from tests.test_html_report import build, resource, quote


class ServiceExportTests(unittest.TestCase):
    def model(self):
        return build([resource('i-a', resource_name='API', instance_type='t3.large'),
                      resource('i-b', resource_name='Worker', instance_type='t3.large'),
                      resource('vol-a', 'ebs', 'volume', storage_type='gp2', storage_gib=100,
                               attached_to=['i-a', 'i-b'], resource_name='Shared data'), quote('vol-a')],
                     language='it', period_start='2026-08-01')

    def test_four_styled_sheets_months_without_dropdowns_and_separate_own_costs(self):
        workbook = load_workbook(BytesIO(build_service_workbook(self.model(), 'ec2')))
        visible = [ws for ws in workbook if ws.sheet_state == 'visible']
        self.assertEqual([ws.title for ws in visible], ['Spesa servizi', 'Per tipo', 'Per risorsa', 'Costi associati'])
        for ws in visible:
            headings = [c.value for c in ws[3]]
            self.assertIn('Regione', headings)
            self.assertIn('2026-08', headings)
            self.assertIn('2026-09', headings)
            self.assertIn('Totale periodo', headings)
            self.assertIsNone(ws.auto_filter.ref)
            self.assertEqual(len(ws.data_validations.dataValidation), 0)
            self.assertEqual(ws.freeze_panes, 'A4')
            self.assertIsNone(ws.sheet_view.pane.xSplit)
            self.assertEqual(ws.sheet_view.pane.ySplit, 3)
            self.assertEqual(ws.sheet_view.pane.topLeftCell, 'A4')
            self.assertEqual(ws.sheet_view.view, 'normal')
            self.assertTrue(ws.sheet_view.showGridLines)
            self.assertEqual(ws['A3'].border.bottom.style, 'thin')
            self.assertTrue(all(not dimension.hidden for dimension in ws.column_dimensions.values()))
            for row in ws:
                for cell in row:
                    if cell.value is not None:
                        self.assertEqual(cell.alignment.horizontal, 'center')
                        self.assertEqual(cell.alignment.vertical, 'center')
            self.assertEqual(ws['A2'].value, 'Account 123456789012')
            self.assertEqual(ws.print_title_rows, '$1:$3')
        own, attached = workbook['Per risorsa'], workbook['Costi associati']
        self.assertEqual(own['F4'].value, 89.28)
        self.assertEqual(attached['E4'].value, 97.28)
        self.assertIn('EBS', attached['D4'].value)
        self.assertEqual(own['F7'].value, '=IF(SUBTOTAL(102,F4:F5)=0,"",SUBTOTAL(109,F4:F5))')
        self.assertEqual(len([row for row in workbook['_unique'] if row[0].value]), 3)
        self.assertIn('_unique', attached['E7'].value)
        self.assertIn('SUBTOTAL(103', workbook['_members']['B1'].value)
        self.assertIn('SUMIF', workbook['_unique']['B1'].value)

    def test_both_download_styles_preserve_all_cell_values_types_and_formulas(self):
        model = self.model()
        exports = embedded_workbooks(model)
        for service, variants in exports.items():
            with self.subTest(service=service):
                self.assertEqual(set(variants), {'basic', 'readable'})
                books = {style: load_workbook(BytesIO(base64.b64decode(export['base64'])))
                         for style, export in variants.items()}
                self.assertNotEqual(variants['basic']['filename'], variants['readable']['filename'])
                self.assertEqual(books['basic'].sheetnames, books['readable'].sheetnames)
                for basic, readable in zip(books['basic'], books['readable']):
                    values = lambda ws: {cell.coordinate: (cell.value, cell.data_type)
                                         for row in ws for cell in row if cell.value is not None}
                    self.assertEqual(values(basic), values(readable), basic.title)
                    self.assertEqual(basic.sheet_state, readable.sheet_state)
                    self.assertEqual((basic.max_row, basic.max_column), (readable.max_row, readable.max_column))
                    self.assertEqual(len(basic.merged_cells.ranges), 0)
                    self.assertIsNone(basic.freeze_panes)
                    self.assertTrue(basic.sheet_view.showGridLines)
                    self.assertTrue(all(not cell.has_style for row in basic for cell in row))
                for book in books.values():
                    book.close()

    def test_basic_download_keeps_formula_like_names_as_text_and_unknown_costs_blank(self):
        model = build([resource('fn', 'lambda', 'function', resource_name='=HYPERLINK("evil")',
                                list_price_per_hour=None)], language='it')
        book = load_workbook(BytesIO(build_service_workbook(model, 'lambda', style='basic')))
        self.assertEqual(book['Per risorsa']['A4'].value, '=HYPERLINK("evil")')
        self.assertEqual(book['Per risorsa']['A4'].data_type, 's')
        self.assertIsNone(book['Per risorsa']['F4'].value)
        self.assertEqual(book['Per risorsa']['G4'].data_type, 'f')

    def test_unknown_costs_are_blank_and_names_are_literal_text(self):
        model = build([resource('fn', 'lambda', 'function', resource_name='=HYPERLINK("evil")',
                                list_price_per_hour=None)], language='it')
        workbook = load_workbook(BytesIO(build_service_workbook(model, 'lambda')))
        self.assertEqual(workbook['Per risorsa']['A4'].data_type, 's')
        self.assertEqual(workbook['Per risorsa']['A4'].value, '=HYPERLINK("evil")')
        self.assertIsNone(workbook['Per risorsa']['F4'].value)
        self.assertIsNone(workbook['Costi associati']['E4'].value)

    def test_billing_only_service_and_native_currency_are_preserved(self):
        bill = {'row_type': 'cost_summary', 'service': 'Special AWS service', 'region': 'us-east-1',
                'period_start': '2026-09-01', 'period_end': '2026-10-01', 'amount': '-5.125', 'unit': 'USD', 'metric': 'AmortizedCost'}
        model = build([bill, {**bill, 'unit': 'EUR', 'amount': '100'}])
        exports = embedded_workbooks(model)
        self.assertIn('billing:Special AWS service', exports)
        workbook = load_workbook(BytesIO(build_service_workbook(model, 'billing:Special AWS service')))
        self.assertEqual(workbook['Service spending']['D4'].value, -5.125)
        self.assertEqual(workbook['Service spending']['A4'].value, 'us-east-1')

    def test_named_typed_links_keep_reference_and_dependency_distinct(self):
        model = build([resource('i-a', resource_name='API', security_group_ids=['sg-a'], related_resource_ids=['vol-a']),
                       resource('vol-a', 'ebs', 'volume', attached_to=['i-a'], storage_gib=100), quote('vol-a'),
                       resource('sg-a', 'vpc', 'security_group', resource_name='Web security', list_price_per_hour=None)])
        volume = next(r for r in model['resources'] if r['id'] == 'vol-a')
        machine = next(r for r in model['resources'] if r['id'] == 'i-a')
        ref = next(r for r in volume['relations'] if r['id'] == 'i-a')
        self.assertEqual((ref['name'], ref['type']), ('API', 'instance'))
        sg = next(r for r in machine['relations'] if r['id'] == 'sg-a')
        self.assertEqual((sg['name'], sg['type']), ('Web security', 'security_group'))
        self.assertNotIn(sg['key'], machine['included_keys'])
        self.assertIn(volume['key'], machine['included_keys'])
        self.assertEqual(machine['totals']['2026-09']['amounts']['USD'], Decimal('94.4'))
