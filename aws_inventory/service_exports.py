"""Styled per-service workbooks embedded in the offline HTML report."""
from __future__ import annotations

import base64
from collections import defaultdict
from decimal import Decimal
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as letter
from openpyxl.workbook.properties import CalcProperties

from .workbooks import _text
from .localization import translate_message

MONEY = '"USD "#,##0.00;[Red]("USD "#,##0.00);"USD "0.00'
HEADER = 3


def known_sum(values):
    known = [Decimal(str(v)) for v in values if v is not None]
    return sum(known, Decimal(0)) if known else None


def own(resource, month):
    value = resource['months'][month]
    return value['amount'] if value['currency'] == 'USD' else None


def service_keys(model):
    keys = set(model['context'].get('inventory_services', [])) | {r['service'] for r in model['resources']}
    mapped = {name for key in keys for name in model['billing_services'].get(key, [])}
    return sorted(keys | {'billing:' + b['service'] for b in model['billing'] if b['service'] not in mapped})


def text_cell(cell, value):
    # Force strings, including strings starting with =, to remain literal Excel text.
    cell.value = _text(value)[:32767]
    cell.data_type = 's'


def build_service_workbook(model, service):
    """Every sheet has calendar-month columns, a region field and totals.

    Attached subtotals can overlap. Hidden membership/unique-cost sheets calculate
    the footer from the union of components of the currently visible parents.
    """
    language = model['context'].get('language', 'en')
    tr = lambda en, italian: translate_message(en, italian, language)
    months = [w['month'] for w in model['windows']]
    resources = [r for r in model['resources'] if r['service'] == service]
    index = {r['key']: r for r in model['resources']}
    names = [service[8:]] if service.startswith('billing:') else model['billing_services'].get(service, [])
    bills = [b for b in model['billing'] if b['service'] in names and b['currency'] == 'USD']
    book = Workbook()
    book.remove(book.active)
    book.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
    book.properties.creator = 'AWS Inventory'
    book.properties.title = service.upper() + ' · USD'

    def sheet(title, columns, rows):
        ws = book.create_sheet(title)
        money_start = len(columns) + 1
        end_col = money_start + len(months)
        if len(rows) > 1_048_576 - HEADER - 2 or end_col > 16384:
            raise ValueError('Service export exceeds Excel worksheet limits.')
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = f'{letter(money_start)}{HEADER + 1}'
        for row, value in [(1, service.upper() + ' · ' + title), (2, f"{tr('Account', 'Account')} {model['account']}")]:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_col)
            cell = ws.cell(row, 1)
            text_cell(cell, value)
            cell.font = Font(name='Calibri', size=20 if row == 1 else 11, bold=row == 1,
                             color='FFFFFF' if row == 1 else '53697B')
            cell.fill = PatternFill('solid', fgColor='17324D' if row == 1 else 'F2F6FA')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            ws.row_dimensions[row].height = 36 if row == 1 else 30
        headers = columns + months + [tr('Period total', 'Totale periodo')]
        for col, value in enumerate(headers, 1):
            cell = ws.cell(HEADER, col)
            text_cell(cell, value)
            cell.fill = PatternFill('solid', fgColor='087F8C')
            cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            ws.column_dimensions[letter(col)].width = 21 if col >= money_start else (30 if col == 1 else 25)
        ws.row_dimensions[HEADER].height = 32
        for rownum, (identity, amounts) in enumerate(rows, HEADER + 1):
            for col, value in enumerate(identity, 1):
                text_cell(ws.cell(rownum, col), value)
            for col, value in enumerate(amounts, money_start):
                if value is not None:
                    ws.cell(rownum, col, Decimal(str(value)))
            start, end = letter(money_start), letter(end_col - 1)
            ws.cell(rownum, end_col, f'=IF(COUNT({start}{rownum}:{end}{rownum})=0,"",SUM({start}{rownum}:{end}{rownum}))')
            for col in range(1, end_col + 1):
                cell = ws.cell(rownum, col)
                cell.font = Font(name='Calibri', size=11, color='243746')
                cell.fill = PatternFill('solid', fgColor='F0F5FA' if rownum % 2 else 'FFFFFF')
                cell.alignment = Alignment(vertical='center', wrap_text=col < money_start)
                if col >= money_start:
                    cell.number_format = MONEY
            ws.row_dimensions[rownum].height = 42
        last = HEADER + len(rows)
        footer = max(HEADER + 2, last + 2)
        text_cell(ws.cell(footer, 1), tr('Total', 'Totale'))
        for col in range(money_start, end_col + 1):
            name = letter(col)
            formula = f'=IF(SUBTOTAL(102,{name}{HEADER + 1}:{name}{last})=0,"",SUBTOTAL(109,{name}{HEADER + 1}:{name}{last}))' if rows else '=""'
            cell = ws.cell(footer, col, formula)
            cell.number_format = MONEY
        for cell in ws[footer][:end_col]:
            cell.fill = PatternFill('solid', fgColor='17324D')
            cell.font = Font(name='Calibri', color='FFFFFF', bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[footer].height = 36
        ws.print_title_rows = f'1:{HEADER}'
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        return ws, money_start, footer

    billed = defaultdict(lambda: defaultdict(list))
    for b in bills:
        billed[(b['region'], b['metric'], b['service'])][b['month']].append(b['amount'])
    sheet(tr('Service spending', 'Spesa servizi'), [tr('Region', 'Regione'), tr('Metric', 'Metrica'), tr('Billing service', 'Servizio fatturato')],
          [(list(key), [known_sum(values[m]) for m in months]) for key, values in sorted(billed.items())])
    grouped = defaultdict(list)
    for r in resources:
        grouped[(r['region'], r['instance_type'] or r['storage_type'] or r['type'])].append(r)
    sheet(tr('By type', 'Per tipo'), [tr('Region', 'Regione'), tr('Instance or resource type', 'Tipo istanza o risorsa'), tr('Coverage', 'Copertura')],
          [([region, kind, tr('Known own costs only', 'Solo costi propri noti')], [known_sum(own(r, m) for r in rs) for m in months]) for (region, kind), rs in sorted(grouped.items())])
    def cost_note(r):
        entries = list(r['months'].values())
        estimated = any(v['basis'] in {'estimated', 'mixed'} for v in entries)
        return tr('Includes estimates', 'Include stime') if estimated else tr('Billed / unavailable', 'Fatturato / non disponibile')
    sheet(tr('By resource', 'Per risorsa'), [tr('Name', 'Nome'), tr('Region', 'Regione'), tr('Type', 'Tipo'), tr('Catalog rate', 'Tariffa catalogo'), tr('Basis', 'Fonte')],
          [([r['name'] or r['id'], r['region'], r['instance_type'] or r['storage_type'] or r['type'], r['plan']['rate_label'] if r['plan']['currency'] == 'USD' else '', cost_note(r)], [own(r, m) for m in months]) for r in resources])
    attached_rows = []
    for r in resources:
        children = [index[k] for k in r['included_keys'] if k != r['key']]
        included = ', '.join(sorted({c['service'].upper() + ' / ' + c['type'] for c in children}))
        included = tr('Includes: ', 'Include: ') + included if included else tr('Own resource only', 'Solo risorsa propria')
        if any(not v['complete'] for v in r['totals'].values()):
            included += tr(' · incomplete known subtotal', ' · subtotale noto incompleto')
        attached_rows.append(([r['name'] or r['id'], r['region'], r['instance_type'] or r['type'], included],
                              [r['totals'][m]['amounts'].get('USD') for m in months]))
    attached, start, footer = sheet(tr('Attached costs', 'Costi associati'), [tr('Name', 'Nome'), tr('Region', 'Regione'), tr('Type', 'Tipo'), tr('Included services', 'Servizi inclusi')], attached_rows)
    attached.column_dimensions['D'].width = 55
    text_cell(attached.cell(footer, 1), tr('Unique total', 'Totale univoco'))
    # Membership rows refer to visible parent cells; unique component rows only
    # contribute if at least one visible parent includes that component.
    membership = book.create_sheet('_members')
    unique = book.create_sheet('_unique')
    membership.sheet_state = unique.sheet_state = 'hidden'
    included_keys = set()
    quoted_title = "'" + attached.title.replace("'", "''") + "'"
    for rownum, r in enumerate(resources, HEADER + 1):
        for key in r['included_keys']:
            included_keys.add(key)
            membership.append([key, f'=SUBTOTAL(103,{quoted_title}!A{rownum})'])
    last_member = max(1, membership.max_row)
    for key in sorted(included_keys):
        rownum = unique.max_row + 1 if unique.cell(1, 1).value is not None else 1
        text_cell(unique.cell(rownum, 1), key)
        for col, m in enumerate(months, 2):
            amount = own(index[key], m)
            if amount is not None:
                unique.cell(rownum, col, f'=IF(SUMIF(_members!$A$1:$A${last_member},A{rownum},_members!$B$1:$B${last_member})>0,{Decimal(str(amount)):f},"")')
    for offset in range(len(months)):
        col = letter(offset + 2)
        rng = f'_unique!{col}1:{col}{max(1, unique.max_row)}'
        attached.cell(footer, start + offset, f'=IF(COUNT({rng})=0,"",SUM({rng}))')
    rng = f'{letter(start)}{footer}:{letter(start + len(months) - 1)}{footer}'
    attached.cell(footer, start + len(months), f'=IF(COUNT({rng})=0,"",SUM({rng}))')
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def embedded_workbooks(model):
    result = {}
    for service in service_keys(model):
        slug = re.sub(r'[^a-zA-Z0-9_-]', '_', service)[:100]
        result[service] = {'filename': f'aws-{ "simulation-" if model["simulation_only"] else ""}{slug}.xlsx',
                           'base64': base64.b64encode(build_service_workbook(model, service)).decode('ascii')}
    return result
