"""Readable service workbooks, companion CSV tables, and staged publication."""

from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import re
import shutil
import tempfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

from .common import compact_json
from .report import cell, redact_sensitive


BUNDLE_FORMAT = "aws-inventory-service-bundle-v1"
HEADER_ROW = 5
MAX_DATA_ROWS = 1_048_576 - HEADER_ROW
NUMERIC_KINDS = {"number", "money", "rate"}
XML_CONTROLS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def _slug(value: str) -> str:
    value = re.sub(r"[^a-z0-9]+", "_", value.lower()).strip("_")
    return value[:90] or "table"


def validate_bundle_destination(path: Path, *, overwrite: bool) -> None:
    """Replace only our own complete bundle, never an arbitrary directory."""
    path = Path(path)
    if path.is_symlink():
        raise ValueError("The report destination cannot be a symbolic link.")
    if not path.exists():
        return
    if not overwrite:
        raise FileExistsError(f"Report already exists: {path}. Use --overwrite.")
    if not path.is_dir():
        raise ValueError("The bundle destination must be a directory.")
    manifest_path = path / "manifest.json"
    if manifest_path.is_symlink() or not manifest_path.is_file():
        raise ValueError("Refusing to replace a directory without a generated report manifest.")
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        if not isinstance(manifest, dict) or manifest.get("format") != BUNDLE_FORMAT:
            raise ValueError("Not a generated service report bundle.")
        files = manifest.get("files")
        if (not isinstance(files, list) or any(not isinstance(name, str) for name in files)
                or len(set(files)) != len(files)):
            raise ValueError("Invalid report manifest file list.")
        expected = set(files) | {"manifest.json"}
        if any(not isinstance(name, str) or Path(name).is_absolute()
               or ".." in Path(name).parts for name in expected):
            raise ValueError("Invalid report manifest paths.")
    except (KeyError, TypeError, json.JSONDecodeError) as exc:
        raise ValueError("Invalid report manifest; refusing replacement.") from exc
    actual = set()
    expected_dirs = {str(parent) for name in expected for parent in Path(name).parents
                     if str(parent) != "."}
    for entry in path.rglob("*"):
        relative = entry.relative_to(path).as_posix()
        if entry.is_symlink():
            raise ValueError(f"Refusing to replace a report containing a symbolic link: {relative}")
        if entry.is_file():
            actual.add(relative)
        elif not entry.is_dir() or relative not in expected_dirs:
            raise ValueError(f"Unexpected entry in report directory: {relative}")
    if actual != expected:
        raise ValueError("The report directory contains missing or additional files; "
                         "choose a new output directory to preserve them.")


def _numeric(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
        raise ValueError(f"Invalid spreadsheet number type: {type(value).__name__}")
    try:
        result = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError("Invalid spreadsheet number.") from exc
    if not result.is_finite():
        raise ValueError("Spreadsheet numbers must be finite.")
    return result


def _text(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple, set)):
        value = compact_json(redact_sensitive(value))
    elif isinstance(value, bool):
        value = "Yes" if value else "No"
    else:
        value = str(value)
    return XML_CONTROLS.sub(lambda match: f"\\u{ord(match.group()):04x}", value)


def _write_csv(path, table, *, excel, simulation=False):
    with path.open("x", encoding="utf-8-sig", newline="") as stream:
        os.chmod(path, 0o600)
        writer = csv.writer(stream, delimiter=";" if excel else ",", lineterminator="\n")
        writer.writerow((["Simulation only"] if simulation else []) +
                        [column.title for column in table.columns])
        for row in table.rows:
            writer.writerow((["true"] if simulation else []) + [
                cell(_numeric(row.get(column.key)), "amount", excel)
                if column.kind in NUMERIC_KINDS else cell(_text(row.get(column.key)), "text", excel)
                for column in table.columns])
        stream.flush()
        os.fsync(stream.fileno())


def _write_workbook(path, report, *, generated_at, context, partial):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.worksheet import Worksheet

    workbook = Workbook(write_only=True)
    workbook.properties.creator = "AWS Inventory"
    workbook.properties.title = report.title
    workbook.properties.description = "Monthly billed costs and current resource inventory"
    navy, teal, pale, orange = "17324D", "087F8C", "F0F5FA", "B75D0A"
    header_font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
    body_font = Font(name="Calibri", size=11, color="243746")
    money_format = '#,##0.00;[Red](#,##0.00);0.00'
    rate_format = '0.000000;[Red](0.000000);0.000000'
    wrap = Alignment(vertical="top", wrap_text=True)
    body_alignment = Alignment(vertical="top", wrap_text=True, indent=1)
    stripe_fills = (PatternFill("solid", fgColor=pale), PatternFill("solid", fgColor="FFFFFF"))
    seen = set()
    try:
        for table in report.tables:
            if len(table.rows) > MAX_DATA_ROWS or len(table.columns) > 16_384:
                raise ValueError("Table exceeds Excel limits; use --format raw-csv for this export.")
            title = re.sub(r"[\\/*?:\[\]]", " ", table.name).strip(" '")[:31] or "Sheet"
            stem, suffix = title, 1
            while title.lower() in seen:
                suffix += 1
                title = stem[:27] + f" {suffix}"
            seen.add(title.lower())
            sheet = workbook.create_sheet(title)
            sheet.freeze_panes = "C6" if len(table.columns) > 3 else "A6"
            sheet.sheet_view.showGridLines = False
            sheet.sheet_view.zoomScale = 90
            financial = any(column.kind == "money" for column in table.columns)
            sheet.sheet_properties.tabColor = orange if financial else teal
            sheet.print_options.horizontalCentered = True
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.paperSize = Worksheet.PAPERSIZE_A3
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.print_title_rows = "1:5"
            sheet.oddFooter.center.text = "Page &P of &N"
            sheet.row_dimensions[1].height = 30
            sheet.row_dimensions[2].height = 24
            sheet.row_dimensions[3].height = 46
            sheet.row_dimensions[HEADER_ROW].height = 32
            banner_end = get_column_letter(max(4, len(table.columns)))
            for banner_row in (1, 2, 3):
                sheet.merged_cells.add(f"A{banner_row}:{banner_end}{banner_row}")
            for index, column in enumerate(table.columns, start=1):
                # Sample only: do not rescan a large table to calculate column widths.
                sample = [len(str(row.get(column.key) or "")) for row in table.rows[:100]]
                width = max([len(column.title) + 2, 14] + sample)
                if column.kind in NUMERIC_KINDS:
                    width = max(15, min(len(column.title) + 2, 24))
                limit = 86 if column.key in {"value", "notes", "price_description"} else 48
                sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, limit)

            def text_cell(value, *, font=body_font, fill=None):
                value = _text(value)
                if len(value) > 32_767:
                    value = value[:32_700] + " [TRUNCATED: full value in companion CSV]"
                result = WriteOnlyCell(sheet, value=value)
                # Force every external string to text, including '=' and URL strings.
                result.data_type = "s"
                result.font, result.alignment = font, wrap
                if fill is not None:
                    result.fill = fill
                return result

            # Banner spans cells visually; table data begins below it at row six.
            sheet.append([text_cell(f"{report.title} | {table.name}",
                                    font=Font(name="Calibri", size=18, color=navy, bold=True))])
            status = "PARTIAL" if partial else "COMPLETE"
            banner = f"{generated_at}  |  Account {context.get('caller_account_id', '')}  |  {status}"
            if context.get("simulation_only"):
                banner = "SIMULATED DATA / NOT AWS BILLING  |  " + banner
            sheet.append([text_cell(banner, font=Font(name="Calibri", size=10, color=orange))])
            sheet.append([text_cell(table.note, font=Font(name="Calibri", size=10, color="536878"))])
            sheet.append([])
            sheet.append([text_cell(column.title, font=header_font,
                                    fill=PatternFill("solid", fgColor=navy))
                          for column in table.columns])
            for row_number, row in enumerate(table.rows, start=HEADER_ROW + 1):
                fill = stripe_fills[row_number % 2]
                # Write-only sheets need row heights set before appending cells.
                line_count = max((sum(max(1, math.ceil(len(line) / max(10,
                    sheet.column_dimensions[get_column_letter(index)].width - 2)))
                    for line in _text(row.get(column.key)).splitlines())
                    for index, column in enumerate(table.columns, start=1)
                    if column.kind not in NUMERIC_KINDS), default=1)
                sheet.row_dimensions[row_number].height = min(180, max(23, line_count * 15))
                values = []
                for column in table.columns:
                    if column.kind in NUMERIC_KINDS:
                        item = WriteOnlyCell(sheet, value=_numeric(row.get(column.key)))
                        item.font, item.fill = body_font, fill
                        item.alignment = Alignment(vertical="top", horizontal="right", indent=1)
                        item.number_format = (money_format if column.kind == "money" else
                                              rate_format if column.kind == "rate" else '#,##0.########')
                    else:
                        item = text_cell(row.get(column.key), fill=fill)
                        item.alignment = body_alignment
                    values.append(item)
                sheet.append(values)
            last_row = HEADER_ROW + len(table.rows)
            sheet.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(table.columns))}{last_row}"
            if financial and table.rows:
                for index, column in enumerate(table.columns, start=1):
                    if column.kind == "money":
                        letter = get_column_letter(index)
                        sheet.conditional_formatting.add(
                            f"{letter}{HEADER_ROW + 1}:{letter}{last_row}",
                            ColorScaleRule(start_type="min", start_color="F7FBFF",
                                           end_type="max", end_color="9BD5CF"))
        preferred = "Monthly by service" if report.key == "overview" else "Monthly by region"
        if preferred in workbook.sheetnames:
            workbook.active = workbook.sheetnames.index(preferred)
        with path.open("xb") as stream:
            os.chmod(path, 0o600)
            workbook.save(stream)
            stream.flush()
            os.fsync(stream.fileno())
    finally:
        # Also clean openpyxl's temporary XML files when serialization fails.
        for sheet in workbook.worksheets:
            writer = getattr(sheet, "_writer", None)
            if writer is not None:
                try:
                    if not sheet.closed:
                        sheet.close()
                    if Path(writer.out).exists():
                        writer.cleanup()
                except (OSError, ValueError):
                    pass
        workbook.close()


def write_service_reports(path: Path, rows: list[dict], *, generated_at: str,
                          context: dict, partial: bool, excel: bool, overwrite: bool):
    """Publish only after every workbook, table and manifest has been written."""
    from .report_views import build_service_reports

    path = path.expanduser().absolute()
    validate_bundle_destination(path, overwrite=overwrite)
    path.parent.mkdir(parents=True, exist_ok=True)
    lock_path = path.parent / f".{path.name}.lock"
    try:
        descriptor = os.open(lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
    except FileExistsError as exc:
        raise FileExistsError(f"Another export owns {lock_path}; choose another destination.") from exc
    os.close(descriptor)
    stage, backup = None, None
    try:
        validate_bundle_destination(path, overwrite=overwrite)
        # Views redact before flattening and skip repeated per-row query contexts.
        reports = build_service_reports(rows, generated_at=generated_at,
                                        context=context, partial=partial)
        stage = Path(tempfile.mkdtemp(prefix=f".{path.name}.staging-", dir=path.parent))
        manifest = {"format": BUNDLE_FORMAT, "generated_at_utc": generated_at,
                    "status": "partial" if partial else "complete", "context": redact_sensitive(context),
                    "simulation_only": bool(context.get("simulation_only")), "files": [], "reports": []}
        names = set()
        for report in reports:
            key = _slug(report.key)
            if key in names:
                raise ValueError(f"Conflicting service filenames: {key}")
            names.add(key)
            workbook_name = ("00_overview" if key == "overview" else key) + ".xlsx"
            _write_workbook(stage / workbook_name, report, generated_at=generated_at,
                            context=context, partial=partial)
            manifest["files"].append(workbook_name)
            tables_dir = stage / "csv" / key
            tables_dir.mkdir(parents=True, mode=0o700)
            table_entries = []
            table_names = set()
            for table in report.tables:
                filename = _slug(table.name)
                if filename in table_names:
                    filename += "_" + hashlib.sha256(table.name.encode()).hexdigest()[:8]
                table_names.add(filename)
                csv_path = tables_dir / (filename + ".csv")
                _write_csv(csv_path, table, excel=excel, simulation=context.get("simulation_only", False))
                relative = csv_path.relative_to(stage).as_posix()
                manifest["files"].append(relative)
                table_entries.append({"name": table.name, "csv": relative,
                                      "rows": len(table.rows), "note": table.note})
            manifest["reports"].append({"service": report.key, "title": report.title,
                                        "workbook": workbook_name, "tables": table_entries})
        with (stage / "manifest.json").open("x", encoding="utf-8") as stream:
            os.chmod(stage / "manifest.json", 0o600)
            stream.write(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")
            stream.flush()
            os.fsync(stream.fileno())
        # Revalidate after staging in case a user edited the prior report meanwhile.
        validate_bundle_destination(path, overwrite=overwrite)
        if path.exists():
            backup = Path(tempfile.mkdtemp(prefix=f".{path.name}.previous-", dir=path.parent))
            backup.rmdir()
            path.rename(backup)
        try:
            stage.rename(path)
            stage = None
        except BaseException:
            if backup is not None:
                backup.rename(path)
                backup = None
            raise
        if backup is not None:
            shutil.rmtree(backup)
            backup = None
    finally:
        if stage is not None:
            shutil.rmtree(stage)
        lock_path.unlink(missing_ok=True)
